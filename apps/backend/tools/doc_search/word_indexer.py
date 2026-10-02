"""
Word Indexer — DOCX document processing

Extracts structured content from Word documents with heading hierarchy.
"""

import sys
import uuid
from pathlib import Path
from typing import Dict, List, Optional

from docx import Document

sys.path.insert(0, str(Path(__file__).parent.parent.parent))

from models.schemas import DocumentChunk, DocumentSearchResult, IndexStatus

from tools.doc_search.base import DocumentIndexer


class WordIndexer(DocumentIndexer):
    """
    Word document indexer for DOCX files.

    Extracts content with heading hierarchy preservation.
    """

    def __init__(self, index_dir: Optional[Path] = None):
        self.index_dir = index_dir or Path(__file__).parent.parent.parent.parent / "data" / "doc_indexes"
        self.index_dir.mkdir(parents=True, exist_ok=True)
        self._status = IndexStatus(source="", source_type="word", chunk_count=0, is_ready=False)

    def ingest(self, source: str, source_type: str = "word") -> List[DocumentChunk]:
        """
        Ingest a DOCX file and return chunks.

        Args:
            source: Path to DOCX file
            source_type: Must be "word"

        Returns:
            List of DocumentChunk objects
        """
        source_path = Path(source)
        if not source_path.exists():
            raise FileNotFoundError(f"DOCX not found: {source}")

        if source_path.suffix.lower() not in [".docx", ".doc"]:
            raise ValueError(f"Not a Word document: {source}")

        # Process DOCX
        result = self._process_docx(str(source_path))

        # Convert to DocumentChunk objects
        chunks = []
        for i, chunk_data in enumerate(result["chunks"]):
            chunk = DocumentChunk(
                content=chunk_data["text"],
                source=str(source_path),
                source_type="word",
                page_number=None,  # DOCX doesn't have pages
                chunk_index=i,
                metadata=chunk_data.get("metadata", {})
            )
            chunks.append(chunk)

        # Update status
        self._status = IndexStatus(
            source=str(source_path),
            source_type="word",
            chunk_count=len(chunks),
            is_ready=len(chunks) > 0
        )

        return chunks

    def search(self, query: str, top_k: int = 5) -> List[DocumentSearchResult]:
        """Search indexed Word documents (requires vector store integration)."""
        raise NotImplementedError("Search requires vector store integration")

    def get_status(self) -> IndexStatus:
        """Get current indexing status."""
        return self._status

    # ========================================================================
    # DOCX PROCESSING (Extracted from mcp/rag.py)
    # ========================================================================

    def _process_docx(self, file_path: str) -> Dict:
        """Process DOCX with structure preservation."""
        sections = self._extract_docx_structure(file_path)
        tables = self._extract_docx_tables(file_path)
        chunks = []

        # Add sections as chunks
        for sec in sections:
            if sec["heading"] == "ROOT" and not sec["content"]:
                continue
            text = f"Section: {sec['heading']}\n"
            for c in sec["content"]:
                text += c["text"] + "\n"
            if text.strip():
                chunks.append({
                    "text": text,
                    "metadata": {
                        "heading": sec["heading"],
                        "level": sec["level"],
                        "type": "section"
                    }
                })

        # Add tables as separate chunks
        for t in tables:
            chunks.append({
                "text": f"Table:\n{t}",
                "metadata": {
                    "type": "table"
                }
            })

        return {"chunks": chunks}

    def _extract_docx_structure(self, file_path: str) -> List[Dict]:
        """Extract DOCX structure with heading hierarchy."""
        doc = Document(file_path)
        sections = []
        current_section = {
            "id": str(uuid.uuid4()),
            "heading": "ROOT",
            "level": 0,
            "content": []
        }

        for para in doc.paragraphs:
            if para.style.name.startswith("Heading"):
                sections.append(current_section)
                level = int(para.style.name.replace("Heading ", ""))
                current_section = {
                    "id": str(uuid.uuid4()),
                    "heading": para.text,
                    "level": level,
                    "content": []
                }
            else:
                if para.text.strip():
                    current_section["content"].append({
                        "type": "paragraph",
                        "text": para.text
                    })

        sections.append(current_section)
        return sections

    def _extract_docx_tables(self, file_path: str) -> List[str]:
        """Extract tables from DOCX as separate chunks."""
        doc = Document(file_path)
        tables = []
        for table in doc.tables:
            table_text = []
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells]
                table_text.append(" | ".join(row_text))
            if table_text:
                tables.append("\n".join(table_text))
        return tables
